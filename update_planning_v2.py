from datetime import datetime, timedelta, date
import json
import os
import re
import time
import zoneinfo
from icalendar import Calendar
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import requests
import urllib3

# Désactivation des avertissements de sécurité SSL si verify=False est utilisé
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ==============================================================================
# 0. GESTION DU FUSEAU HORAIRE (EUROPE/PARIS)
# ==============================================================================

TZ_PARIS = zoneinfo.ZoneInfo("Europe/Paris")


def convertir_en_heure_paris(dt_object):
    """Gère l'attribution du fuseau horaire Europe/Paris sans appliquer
    de décalage UTC intempestif sur les heures locales transmises par ADE.
    """
    if dt_object is None:
        return None

    if hasattr(dt_object, "dt"):
        dt_object = dt_object.dt

    if not isinstance(dt_object, datetime):
        return dt_object

    if dt_object.tzinfo is None:
        return dt_object.replace(tzinfo=TZ_PARIS)

    return dt_object.astimezone(TZ_PARIS)


# ==============================================================================
# 1. CONFIGURATION
# ==============================================================================

URLS_ADE = {
    "LP MIA": (
        "https://agenda-web-consult.univ-amu.fr/jsp/custom/modules/plannings/anonymous_cal.jsp?projectId=8&resources=127685,127687,127688&calType=ical&firstDate=2026-08-17&lastDate=2027-08-15"
    ),
    "BUT 3 SNRV": (
        "https://agenda-web-consult.univ-amu.fr/jsp/custom/modules/plannings/anonymous_cal.jsp?projectId=8&resources=58374,58375&calType=ical&firstDate=2026-08-17&lastDate=2027-08-15"
    ),
    "M1 MPAD": (
        "https://agenda-web-consult.univ-amu.fr/jsp/custom/modules/plannings/anonymous_cal.jsp?projectId=8&resources=70202&calType=ical&firstDate=2026-08-17&lastDate=2027-08-15"
    ),
    "M2 MPAD": (
        "https://agenda-web-consult.univ-amu.fr/jsp/custom/modules/plannings/anonymous_cal.jsp?projectId=8&resources=391&calType=ical&firstDate=2026-08-17&lastDate=2027-08-15"
    ),
}

ENSEIGNANTS_AUTORISES = [
    "ATTAFI",
    "CORNUEAU",
    "GUEUDRE",
    "VALLEE",
    "SANCHEZ",
    "CHAVES-JACOB",
    "AMADEI",
    "MAZOYER",
    "RAYNAL",
    "MOYSAN",
]

FICHIER_EXCEL_SORTIE = "Planning_Voitures_Tallard.xlsx"
FICHIER_EXCEL_TEMP = "Planning_Voitures_Tallard_temp.xlsx"
FICHIER_HTML_SORTIE = "Planning_Voitures_Tallard.html"

JOURS_FR = {
    "Monday": "Lundi",
    "Tuesday": "Mardi",
    "Wednesday": "Mercredi",
    "Thursday": "Jeudi",
    "Friday": "Vendredi",
    "Saturday": "Samedi",
    "Sunday": "Dimanche",
}


# ==============================================================================
# 2. EXTRACTION ADE
# ==============================================================================


def telecharger_ical_avec_retry(url, retries=3, backoff_factor=1):
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
    }
    for i in range(retries):
        try:
            response = requests.get(url, headers=headers, timeout=(5, 5))
            response.raise_for_status()
            return response.content
        except requests.RequestException as e:
            print(f"⚠️ Tentative {i + 1}/{retries} échouée")
            if i == retries - 1:
                raise e
            time.sleep(backoff_factor * (i + 1))


def est_enseignant_autorise(texte_evenement):
    texte_lower = texte_evenement.lower()
    for nom in ENSEIGNANTS_AUTORISES:
        if nom.lower() in texte_lower:
            return True, nom.upper()
    return False, "Vacataire / Inconnu"


def est_en_visio(texte_evenement):
    return bool(re.search(r"visio salle", texte_evenement, re.IGNORECASE))


def extraire_cours():
    tous_les_cours = []
    aujourdhui = date.today()
    il_y_a_2j = aujourdhui - timedelta(days=2)
    dans_7j = aujourdhui + timedelta(days=7)

    for cohorte, url in URLS_ADE.items():
        print(f"📥 Téléchargement du flux ADE pour {cohorte}...")
        try:
            contenu_ical = telecharger_ical_avec_retry(url)
            cal = Calendar.from_ical(contenu_ical)

            for component in cal.walk("VEVENT"):
                summary = str(component.get("summary", ""))
                description = str(component.get("description", ""))
                location = str(component.get("location", ""))

                dtstart_raw = component.get("dtstart")
                dtend_raw = component.get("dtend")

                if not dtstart_raw:
                    continue

                dtstart = convertir_en_heure_paris(dtstart_raw)
                dtend = convertir_en_heure_paris(dtend_raw)

                if not isinstance(dtstart, datetime):
                    continue

                texte_complet = f"{summary} {description} {location}"
                est_autorise, nom_prof = est_enseignant_autorise(texte_complet)
                visio_detectee = est_en_visio(texte_complet)

                if not est_autorise:
                    statut_voiture = "NON (Vacataire / Perso)"
                    code_statut = "NON"
                    motif = "Enseignant non habilité voiture de service"
                elif visio_detectee:
                    statut_voiture = "NON (Visio salle)"
                    code_statut = "NON"
                    motif = "Cours assuré à distance"
                else:
                    statut_voiture = "OUI (Voiture requise)"
                    code_statut = "OUI"
                    motif = "Présentiel - Enseignant habilité"

                # On ne conserve que les cours nécessitant une voiture de service
                if code_statut != "OUI":
                    continue

                dt_date = dtstart.date()
                jour_fr = JOURS_FR.get(
                    dtstart.strftime("%A"), dtstart.strftime("%A")
                )

                est_2j_passes = il_y_a_2j <= dt_date < aujourdhui
                est_7j_futurs = aujourdhui <= dt_date <= dans_7j
                est_vieux_passe = dt_date < il_y_a_2j

                if est_2j_passes:
                    ordre_groupe = 0
                elif est_7j_futurs:
                    ordre_groupe = 1
                elif dt_date > dans_7j:
                    ordre_groupe = 2
                else:
                    ordre_groupe = 3

                tous_les_cours.append({
                    "id": f"{dtstart.strftime('%Y%m%d%H%M')}_{cohorte}_{nom_prof}",
                    "_dt_start": dtstart,
                    "_dt_end": dtend,
                    "date_iso": dt_date.isoformat(),
                    "Date": dtstart.strftime("%d/%m/%Y"),
                    "Jour": jour_fr,
                    "Heure Début": dtstart.strftime("%H:%M"),
                    "Heure Fin": dtend.strftime("%H:%M"),
                    "Horaires": (
                        f"{dtstart.strftime('%H:%M')} -"
                        f" {dtend.strftime('%H:%M')}"
                    ),
                    "Cohorte": cohorte,
                    "Matière / Cours": summary,
                    "Enseignant détecté": nom_prof,
                    "Salle / Équipement": location,
                    "Besoin Voiture Service": statut_voiture,
                    "code_statut": code_statut,
                    "Explication": motif,
                    "est_7j_passes": est_2j_passes,
                    "est_7j_futurs": est_7j_futurs,
                    "est_vieux_passe": est_vieux_passe,
                    "ordre_groupe": ordre_groupe,
                    "start_minutes": dtstart.hour * 60 + dtstart.minute,
                    "end_minutes": dtend.hour * 60 + dtend.minute,
                })

        except Exception as e:
            print(f"❌ Échec de la récupération pour {cohorte}: {e}")

    tous_les_cours.sort(key=lambda x: (x["ordre_groupe"], x["_dt_start"]))
    return tous_les_cours


# ==============================================================================
# 3. GENERATION EXCEL V2
# ==============================================================================


def appliquer_mise_en_forme_excel(chemin_fichier):
    wb = openpyxl.load_workbook(chemin_fichier)

    fill_header = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    fill_bleu = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )
    font_bleu = Font(name="Calibri", size=10, color="1F4E78", bold=True)

    fill_vert = PatternFill(
        start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
    )
    font_vert = Font(name="Calibri", size=10, color="375623", bold=True)

    fill_rouge = PatternFill(
        start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
    )
    font_rouge = Font(name="Calibri", size=10, color="C00000")

    font_normal = Font(name="Calibri", size=10)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    aujourdhui = date.today()
    il_y_a_7j = aujourdhui - timedelta(days=7)
    dans_7j = aujourdhui + timedelta(days=7)

    for sheetname in wb.sheetnames:
        ws = wb[sheetname]

        for cell in ws[1]:
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(
                min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column
        ):
            date_str = ws.cell(row=row[0].row, column=1).value
            dt_cours = None
            if date_str:
                try:
                    dt_cours = datetime.strptime(
                        str(date_str), "%d/%m/%Y"
                    ).date()
                except ValueError:
                    pass

            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                if dt_cours:
                    if il_y_a_7j <= dt_cours < aujourdhui:
                        cell.fill = fill_bleu
                        cell.font = font_bleu
                    elif aujourdhui <= dt_cours <= dans_7j:
                        cell.fill = fill_vert
                        cell.font = font_vert
                    elif dt_cours < il_y_a_7j:
                        cell.fill = fill_rouge
                        cell.font = font_rouge
                    else:
                        cell.font = font_normal

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(chemin_fichier)


def generer_excel(liste_cours):
    if not liste_cours:
        return

    df = pd.DataFrame(liste_cours)
    cols_to_drop = [
        "id",
        "_dt_start",
        "_dt_end",
        "date_iso",
        "Horaires",
        "code_statut",
        "est_7j_passes",
        "est_7j_futurs",
        "est_vieux_passe",
        "ordre_groupe",
        "start_minutes",
        "end_minutes",
    ]
    df_clean = df.drop(columns=[c for c in cols_to_drop if c in df.columns])

    with pd.ExcelWriter(FICHIER_EXCEL_TEMP, engine="openpyxl") as writer:
        df_clean.to_excel(
            writer, sheet_name="Voitures à réserver", index=False
        )

    appliquer_mise_en_forme_excel(FICHIER_EXCEL_TEMP)

    for _ in range(5):
        try:
            if os.path.exists(FICHIER_EXCEL_SORTIE):
                os.remove(FICHIER_EXCEL_SORTIE)
            os.rename(FICHIER_EXCEL_TEMP, FICHIER_EXCEL_SORTIE)
            print(
                "📊 Fichier Excel v2 généré :"
                f" {os.path.abspath(FICHIER_EXCEL_SORTIE)}"
            )
            break
        except PermissionError:
            time.sleep(3)


# ==============================================================================
# 4. GENERATION HTML V2 (AVEC OPTION 1 CORRIGÉE & OPTION 2 INTACTE)
# ==============================================================================


def generer_tableau_html_option2(cours_list):
    if not cours_list:
        return (
            "<p class='p-6 text-center text-slate-500 font-medium'>Aucune"
            " réservation de voiture nécessaire.</p>"
        )

    html = """
    <div class="overflow-x-auto">
        <table class="w-full text-left border-collapse text-sm">
            <thead>
                <tr class="bg-slate-800 text-white uppercase text-xs tracking-wider">
                    <th class="p-3">Date</th>
                    <th class="p-3">Jour</th>
                    <th class="p-3">Horaire</th>
                    <th class="p-3">Cohorte</th>
                    <th class="p-3">Matière / Cours</th>
                    <th class="p-3">Enseignant</th>
                    <th class="p-3">Salle</th>
                </tr>
            </thead>
            <tbody class="divide-y divide-slate-200">
    """
    for c in cours_list:
        if c["est_7j_passes"]:
            row_class = "bg-blue-50 text-blue-900 font-medium"
        elif c["est_7j_futurs"]:
            row_class = "bg-emerald-50 text-emerald-900 font-medium"
        elif c["est_vieux_passe"]:
            row_class = "bg-red-50 text-red-700 opacity-75"
        else:
            row_class = "bg-white text-slate-800 hover:bg-slate-50"

        html += f"""
        <tr class="{row_class} transition-colors" data-search-text="{c['Date']} {c['Jour']} {c['Cohorte']} {c['Matière / Cours']} {c['Enseignant détecté']} {c['Salle / Équipement']}">
            <td class="p-3 font-semibold">{c['Date']}</td>
            <td class="p-3">{c['Jour']}</td>
            <td class="p-3 whitespace-nowrap">{c['Horaires']}</td>
            <td class="p-3"><span class="px-2 py-1 rounded bg-slate-200 text-slate-800 text-xs font-bold">{c['Cohorte']}</span></td>
            <td class="p-3 font-medium">{c['Matière / Cours']}</td>
            <td class="p-3">{c['Enseignant détecté']}</td>
            <td class="p-3">{c['Salle / Équipement']}</td>
        </tr>
        """
    html += "</tbody></table></div>"
    return html


def generer_html_v2(cours):
    date_maj = datetime.now(TZ_PARIS).strftime("%d/%m/%Y à %H:%M")

    cours_js_data = []
    for c in cours:
        cours_js_data.append({
            "id": c["id"],
            "date": c["date_iso"],
            "date_fr": c["Date"],
            "jour": c["Jour"],
            "start_time": c["Heure Début"],
            "end_time": c["Heure Fin"],
            "horaires": c["Horaires"],
            "cohorte": c["Cohorte"],
            "cours": c["Matière / Cours"],
            "enseignant": c["Enseignant détecté"],
            "salle": c["Salle / Équipement"],
            "start_minutes": c["start_minutes"],
            "end_minutes": c["end_minutes"]
        })

    json_cours_str = json.dumps(cours_js_data, ensure_ascii=False)
    tableau_option2_html = generer_tableau_html_option2(cours)

    html_content = """<!DOCTYPE html>
<html lang="fr" class="h-full bg-slate-100">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Planning Voiture — Tallard v2</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <style>
        /* Force l'espace de la scrollbar sur les conteneurs pour aligner parfaitement la grille */
        .sync-scroll-container {
            overflow-y: scroll;
        }
        .grid-layout {
            display: grid;
            grid-template-columns: 60px repeat(5, minmax(0, 1fr));
            width: 100%;
        }
        .time-line {
            position: absolute;
            left: 60px;
            right: 0;
            border-top: 1.5px dashed #cbd5e1;
            pointer-events: none;
            z-index: 10;
        }
        .time-label {
            position: absolute;
            left: 4px;
            transform: translateY(-50%);
            font-size: 10px;
            font-family: monospace;
            color: #64748b;
            background-color: #f8fafc;
            padding: 0 2px;
            border-radius: 2px;
            z-index: 11;
        }
    </style>
</head>
<body class="h-full font-sans antialiased text-slate-900 flex flex-col">

    <!-- BARRE SUPÉRIEURE FIXE -->
    <header class="bg-slate-900 text-white p-4 shadow-md shrink-0">
        <div class="max-w-7xl mx-auto flex flex-col md:flex-row items-stretch md:items-center justify-between gap-4">

            <div class="flex items-center justify-between md:justify-start gap-4">
                <div>
                    <h1 class="text-xl font-bold flex items-center gap-2">
                        <span>🚗</span>
                        <span>Planning Voiture — Tallard</span>
                        <span class="text-xs font-normal bg-slate-800 text-slate-300 px-2 py-0.5 rounded-full">v2</span>
                    </h1>
                    <p class="text-slate-400 text-xs mt-0.5">Dernière synchronisation ADE : <span class="text-slate-200 font-semibold">__DATE_MAJ__</span></p>
                </div>
            </div>

            <div class="flex items-center gap-3 self-end md:self-auto">
                <div class="relative w-64">
                    <input type="text" id="searchInput" oninput="handleSearch()" placeholder="🔍 Filtrer (Nom, Promo, Salle)..." 
                           class="w-full bg-slate-800 text-slate-100 placeholder-slate-400 text-xs px-3 py-2 rounded-lg border border-slate-700 focus:outline-none focus:ring-2 focus:ring-blue-500">
                </div>

                <div class="bg-slate-800 p-1 rounded-lg flex gap-1 border border-slate-700 shrink-0">
                    <button id="btnViewGrid" onclick="switchView('grid')" class="px-3 py-1.5 text-xs font-medium rounded-md transition-all bg-blue-600 text-white shadow">
                        📅 Emploi du temps
                    </button>
                    <button id="btnViewTable" onclick="switchView('table')" class="px-3 py-1.5 text-xs font-medium rounded-md transition-all text-slate-300 hover:text-white hover:bg-slate-700">
                        📋 Vue Tableau
                    </button>
                </div>
            </div>
        </div>
    </header>

    <!-- CONTENU PRINCIPAL -->
    <main class="flex-1 max-w-7xl w-full mx-auto p-4 flex flex-col overflow-hidden">

        <!-- BARRE DE LÉGENDE CONDITIONNELLE -->
        <div class="bg-white p-3 rounded-lg shadow-sm border border-slate-200 mb-4 flex items-center justify-between text-xs shrink-0">
            <!-- Légende pour l'emploi du temps -->
            <div id="gridLegend" class="flex items-center gap-3 flex-wrap">
                <span class="font-semibold text-slate-500">Cohortes:</span>
                <span class="flex items-center gap-1.5 px-2 py-1 rounded bg-blue-100 text-blue-900 border border-blue-200 font-semibold"><span class="w-2.5 h-2.5 rounded-full bg-blue-600"></span> LP MIA</span>
                <span class="flex items-center gap-1.5 px-2 py-1 rounded bg-emerald-100 text-emerald-900 border border-emerald-200 font-semibold"><span class="w-2.5 h-2.5 rounded-full bg-emerald-600"></span> BUT 3 SNRV</span>
                <span class="flex items-center gap-1.5 px-2 py-1 rounded bg-amber-100 text-amber-900 border border-amber-200 font-semibold"><span class="w-2.5 h-2.5 rounded-full bg-amber-600"></span> M1 MPAD</span>
                <span class="flex items-center gap-1.5 px-2 py-1 rounded bg-purple-100 text-purple-900 border border-purple-200 font-semibold"><span class="w-2.5 h-2.5 rounded-full bg-purple-600"></span> M2 MPAD</span>
            </div>

            <!-- Légende pour la vue tableau (texte "Statut réservation :" supprimé) -->
            <div id="tableLegend" class="hidden flex items-center gap-4 text-slate-600 font-medium">
                <span class="flex items-center gap-1.5 px-2 py-0.5 rounded bg-blue-100 text-blue-900"><span class="w-2.5 h-2.5 rounded-full bg-blue-500"></span> 2 jours passés</span>
                <span class="flex items-center gap-1.5 px-2 py-0.5 rounded bg-emerald-100 text-emerald-900"><span class="w-2.5 h-2.5 rounded-full bg-emerald-500"></span> 7 jours futurs</span>
                <span class="flex items-center gap-1.5 px-2 py-0.5 rounded bg-red-100 text-red-800"><span class="w-2.5 h-2.5 rounded-full bg-red-400"></span> Ancien (-2j)</span>
            </div>
        </div>

        <!-- VIEW 1 : EMPLOI DU TEMPS -->
        <div id="viewGrid" class="flex-1 bg-white rounded-lg shadow-sm border border-slate-200 flex flex-col min-h-0 overflow-hidden">

            <div class="p-3 border-b border-slate-200 bg-slate-50 grid grid-cols-3 items-center shrink-0">
                <div></div>
                <div class="flex items-center justify-center gap-1">
                    <button onclick="changeWeek(-1)" class="p-1.5 hover:bg-slate-200 rounded-md text-slate-600 transition" title="Semaine précédente">
                        ◀
                    </button>
                    <button onclick="goToday()" class="px-3 py-1 bg-slate-200 hover:bg-slate-300 text-slate-700 font-semibold text-xs rounded-md transition">
                        Aujourd'hui
                    </button>
                    <button onclick="changeWeek(1)" class="p-1.5 hover:bg-slate-200 rounded-md text-slate-600 transition" title="Semaine suivante">
                        ▶
                    </button>
                </div>
                <h2 id="currentWeekTitle" class="text-xs font-bold text-slate-800 text-right">
                    --
                </h2>
            </div>

            <!-- EN-TÊTE DES JOURS (MÊME STRUCTURE DE DEFILEMENT QUE LE CORPS POUR ALIGNEMENT STABILISÉ) -->
            <div class="sync-scroll-container overflow-y-scroll border-b border-slate-200 bg-slate-100 text-center font-bold text-xs text-slate-700 shrink-0">
                <div class="grid-layout">
                    <div class="p-2 border-r border-slate-200 flex items-center justify-center text-slate-400">H</div>
                    <div id="day-header-0" class="p-2 border-r border-slate-200">Lundi</div>
                    <div id="day-header-1" class="p-2 border-r border-slate-200">Mardi</div>
                    <div id="day-header-2" class="p-2 border-r border-slate-200">Mercredi</div>
                    <div id="day-header-3" class="p-2 border-r border-slate-200">Jeudi</div>
                    <div id="day-header-4" class="p-2">Vendredi</div>
                </div>
            </div>

            <!-- CORPS DU TABLEAU DE BORD -->
            <div class="flex-1 sync-scroll-container overflow-y-scroll relative min-h-0">
                <div class="grid-layout relative h-[720px]">

                    <!-- LIGNES HORIZONTAUX DES CRÉNEAUX SPÉCIFIQUES -->
                    <div class="time-line" style="top: 90px;"></div>
                    <div class="time-label" style="top: 90px;">08:30</div>

                    <div class="time-line" style="top: 210px;"></div>
                    <div class="time-label" style="top: 210px;">10:30</div>

                    <div class="time-line" style="top: 330px; border-top-color: #94a3b8;"></div>
                    <div class="time-label" style="top: 330px; font-weight: bold; color: #475569;">12:30</div>

                    <div class="time-line" style="top: 405px;"></div>
                    <div class="time-label" style="top: 405px;">13:45</div>

                    <div class="time-line" style="top: 525px;"></div>
                    <div class="time-label" style="top: 525px;">15:45</div>

                    <div class="time-line" style="top: 645px;"></div>
                    <div class="time-label" style="top: 645px;">17:45</div>

                    <!-- COLONNE HEURES DE BASE -->
                    <div class="border-r border-slate-200 bg-slate-50 text-[11px] font-mono text-slate-400 flex flex-col justify-between py-1 select-none">
                        <div class="h-[60px] text-center">07:00</div>
                        <div class="h-[60px] text-center">08:00</div>
                        <div class="h-[60px] text-center">09:00</div>
                        <div class="h-[60px] text-center">10:00</div>
                        <div class="h-[60px] text-center">11:00</div>
                        <div class="h-[60px] text-center">12:00</div>
                        <div class="h-[60px] text-center">13:00</div>
                        <div class="h-[60px] text-center">14:00</div>
                        <div class="h-[60px] text-center">15:00</div>
                        <div class="h-[60px] text-center">16:00</div>
                        <div class="h-[60px] text-center">17:00</div>
                        <div class="h-[60px] text-center">18:00</div>
                        <div class="h-[60px] text-center">19:00</div>
                    </div>

                    <!-- COLONNES ALIGNÉES -->
                    <div id="col-day-0" class="border-r border-slate-200 relative bg-white h-full"></div>
                    <div id="col-day-1" class="border-r border-slate-200 relative bg-white h-full"></div>
                    <div id="col-day-2" class="border-r border-slate-200 relative bg-white h-full"></div>
                    <div id="col-day-3" class="border-r border-slate-200 relative bg-white h-full"></div>
                    <div id="col-day-4" class="relative bg-white h-full"></div>

                </div>
            </div>
        </div>

        <!-- VIEW 2 : TABLEAU DÉTAILLÉ -->
        <div id="viewTable" class="hidden flex-1 bg-white rounded-lg shadow-sm border border-slate-200 overflow-y-auto">
            __TABLEAU_OPTION_2__
        </div>

    </main>

    <!-- JS FRONTEND -->
    <script>
        const COURS_DATA = __JSON_COURS_DATA__;
        let currentMonday = getMonday(new Date());
        let currentView = 'grid';

       function getMonday(d) {
            d = new Date(d);
            const day = d.getDay(); // 0: Dimanche, 1: Lundi, ..., 6: Samedi
            
            // Si c'est Samedi (6) ou Dimanche (0), on bascule directement sur la semaine suivante
            if (day === 6 || day === 0) {
                const daysToAdd = (day === 6) ? 2 : 1;
                d.setDate(d.getDate() + daysToAdd);
            } else {
                // Du lundi (1) au vendredi (5) : on recule du nombre de jours nécessaires depuis le lundi
                const diffToMonday = day - 1; // Lundi -> 0, Mardi -> 1, ..., Vendredi -> 4
                d.setDate(d.getDate() - diffToMonday);
            }
        
            d.setHours(0, 0, 0, 0);
            return d;
        }

        function formatDateFR(dateObj) {
            const d = String(dateObj.getDate()).padStart(2, '0');
            const m = String(dateObj.getMonth() + 1).padStart(2, '0');
            return `${d}/${m}/${dateObj.getFullYear()}`;
        }

        function switchView(view) {
            currentView = view;
            const btnGrid = document.getElementById('btnViewGrid');
            const btnTable = document.getElementById('btnViewTable');
            const viewGrid = document.getElementById('viewGrid');
            const viewTable = document.getElementById('viewTable');
            const gridLegend = document.getElementById('gridLegend');
            const tableLegend = document.getElementById('tableLegend');

            if (view === 'grid') {
                viewGrid.classList.remove('hidden');
                viewTable.classList.add('hidden');
                gridLegend.classList.remove('hidden');
                tableLegend.classList.add('hidden');

                btnGrid.className = 'px-3 py-1.5 text-xs font-medium rounded-md transition-all bg-blue-600 text-white shadow';
                btnTable.className = 'px-3 py-1.5 text-xs font-medium rounded-md transition-all text-slate-300 hover:text-white hover:bg-slate-700';
                renderGrid();
            } else {
                viewGrid.classList.add('hidden');
                viewTable.classList.remove('hidden');
                gridLegend.classList.add('hidden');
                tableLegend.classList.remove('hidden');

                btnTable.className = 'px-3 py-1.5 text-xs font-medium rounded-md transition-all bg-blue-600 text-white shadow';
                btnGrid.className = 'px-3 py-1.5 text-xs font-medium rounded-md transition-all text-slate-300 hover:text-white hover:bg-slate-700';
            }
        }

        function changeWeek(direction) {
            currentMonday.setDate(currentMonday.getDate() + (direction * 7));
            renderGrid();
        }

        function goToday() {
            currentMonday = getMonday(new Date());
            renderGrid();
        }

        function getCohorteColor(cohorte) {
            switch(cohorte) {
                case 'LP MIA': return 'bg-blue-100 text-blue-900 border-blue-300 hover:bg-blue-200';
                case 'BUT 3 SNRV': return 'bg-emerald-100 text-emerald-900 border-emerald-300 hover:bg-emerald-200';
                case 'M1 MPAD': return 'bg-amber-100 text-amber-900 border-amber-300 hover:bg-amber-200';
                case 'M2 MPAD': return 'bg-purple-100 text-purple-900 border-purple-300 hover:bg-purple-200';
                default: return 'bg-slate-100 text-slate-900 border-slate-300 hover:bg-slate-200';
            }
        }

        function renderGrid() {
            const searchFilter = document.getElementById('searchInput').value.toLowerCase().trim();
            const daysNames = ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi'];
            const weekDates = [];

            for (let i = 0; i < 5; i++) {
                const d = new Date(currentMonday);
                d.setDate(d.getDate() + i);
                weekDates.push(d);

                const headerEl = document.getElementById(`day-header-${i}`);
                const isToday = new Date().toDateString() === d.toDateString();
                headerEl.innerHTML = `
                    <div class="${isToday ? 'text-blue-600 font-extrabold' : ''}">${daysNames[i]}</div>
                    <div class="text-[10px] font-normal text-slate-500">${formatDateFR(d)}</div>
                `;
            }

            const sunday = new Date(currentMonday);
            sunday.setDate(sunday.getDate() + 4);
            document.getElementById('currentWeekTitle').innerText = 
                `Semaine du ${formatDateFR(currentMonday)} au ${formatDateFR(sunday)}`;

            for (let i = 0; i < 5; i++) {
                const col = document.getElementById(`col-day-${i}`);
                col.innerHTML = '';
            }

            const startWeekIso = currentMonday.toISOString().split('T')[0];
            const endWeekIso = sunday.toISOString().split('T')[0];

            const weekEvents = COURS_DATA.filter(c => {
                const matchesWeek = c.date >= startWeekIso && c.date <= endWeekIso;
                if (!matchesWeek) return false;

                if (searchFilter) {
                    const fullText = `${c.enseignant} ${c.cohorte} ${c.cours} ${c.salle}`.toLowerCase();
                    return fullText.includes(searchFilter);
                }
                return true;
            });

            for (let dayIdx = 0; dayIdx < 5; dayIdx++) {
                const dayDateStr = weekDates[dayIdx].toISOString().split('T')[0];
                const dayEvents = weekEvents.filter(e => e.date === dayDateStr);
                const colEl = document.getElementById(`col-day-${dayIdx}`);

                if (dayEvents.length === 0) continue;

                dayEvents.sort((a, b) => a.start_minutes - b.start_minutes);

                const columns = [];
                dayEvents.forEach(event => {
                    let placed = false;
                    for (let c = 0; c < columns.length; c++) {
                        const lastEvent = columns[c][columns[c].length - 1];
                        if (lastEvent.end_minutes <= event.start_minutes) {
                            columns[c].push(event);
                            event.colIdx = c;
                            placed = true;
                            break;
                        }
                    }
                    if (!placed) {
                        event.colIdx = columns.length;
                        columns.push([event]);
                    }
                });

                const totalCols = columns.length;

                dayEvents.forEach(event => {
                    const topPx = ((event.start_minutes - 420) / 60) * 60;
                    const heightPx = Math.max(((event.end_minutes - event.start_minutes) / 60) * 60, 24);
                    const widthPercent = 100 / totalCols;
                    const leftPercent = event.colIdx * widthPercent;

                    const colorStyle = getCohorteColor(event.cohorte);

                    const card = document.createElement('div');
                    card.className = `absolute rounded p-1 border text-xs shadow-sm overflow-hidden flex flex-col justify-center leading-tight transition-all cursor-pointer z-20 ${colorStyle}`;
                    card.style.top = `${topPx}px`;
                    card.style.height = `${heightPx}px`;
                    card.style.width = `calc(${widthPercent}% - 2px)`;
                    card.style.left = `${leftPercent}%`;

                    card.innerHTML = `
                        <div class="font-bold text-[11px] truncate uppercase tracking-wide" title="${event.enseignant}">${event.enseignant}</div>
                        ${heightPx >= 40 ? `<div class="text-[10px] opacity-90 truncate font-medium">${event.cohorte}</div>` : ''}
                    `;

                    card.title = `${event.enseignant} (${event.cohorte})\nHoraires: ${event.horaires}\nCours: ${event.cours}\nSalle: ${event.salle}`;

                    colEl.appendChild(card);
                });
            }
        }

        function handleSearch() {
            if (currentView === 'grid') {
                renderGrid();
            } else {
                const input = document.getElementById('searchInput').value.toLowerCase();
                const rows = document.querySelectorAll('#viewTable tbody tr');
                rows.forEach(row => {
                    const text = row.getAttribute('data-search-text').toLowerCase();
                    row.style.display = text.includes(input) ? '' : 'none';
                });
            }
        }

        document.addEventListener('DOMContentLoaded', () => {
            renderGrid();
        });
    </script>
</body>
</html>
"""

    html_content = html_content.replace("__DATE_MAJ__", date_maj)
    html_content = html_content.replace("__TABLEAU_OPTION_2__", tableau_option2_html)
    html_content = html_content.replace("__JSON_COURS_DATA__", json_cours_str)

    with open(FICHIER_HTML_SORTIE, "w", encoding="utf-8") as f:
        f.write(html_content)

    print(f"🌐 Fichier HTML v2 généré avec succès : {os.path.abspath(FICHIER_HTML_SORTIE)}")
          
 # ==============================================================================
# 5. EXECUTION
    #==============================================================================
          
if __name__ == "__main__":
    cours = extraire_cours()

# Génération des fichiers de sortie
    generer_html_v2(cours)

if cours:
  generer_excel(cours)
  print(
      "\n✅ Traitement V2 terminé : Fichiers HTML et Excel générés avec"
      " succès."
  )
